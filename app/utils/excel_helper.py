import io
import logging
import os
from pathlib import Path
from typing import Optional
from PIL import Image as PILImage
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Border, Side, Font, PatternFill, Alignment

logger = logging.getLogger(__name__)

# Preferred paths to discover the corporate logo
DEFAULT_LOGO_CANDIDATES = [
    os.getenv("EXCEL_LOGO_PATH"),
    r"C:\Users\User\Downloads\abs-nav-footer-light.webp",
    str(Path(__file__).resolve().parent.parent / "assets" / "logo.png"),
    str(Path(__file__).resolve().parent.parent / "assets" / "abs-nav-footer-light.webp"),
]


def get_report_logo(target_height: int = 54) -> Optional[OpenpyxlImage]:
    """
    Locates, processes, and formats the corporate logo for Openpyxl.
    Handles .webp transparency, crops empty border padding, and converts
    to in-memory PNG for full OpenXML Excel compatibility.
    """
    for candidate in DEFAULT_LOGO_CANDIDATES:
        if not candidate or not os.path.exists(candidate):
            continue
        try:
            pil_img = PILImage.open(candidate).convert("RGBA")
            # Auto-crop transparent boundaries for crisp alignment
            bbox = pil_img.getbbox()
            if bbox:
                pil_img = pil_img.crop(bbox)

            buf = io.BytesIO()
            pil_img.save(buf, format="PNG")
            buf.seek(0)

            img = OpenpyxlImage(buf)
            aspect_ratio = img.width / max(img.height, 1)
            img.height = target_height
            img.width = int(target_height * aspect_ratio)
            return img
        except Exception as e:
            logger.warning("Failed to load Excel report logo from %s: %s", candidate, e)

    return None


def center_image_in_cell(
    ws,
    img: OpenpyxlImage,
    col_idx: int = 0,
    row_idx: int = 1,
    col_width_chars: float = 18,
    row_offset_px: int = 3,
):
    """
    Centers an image horizontally within a given column and anchors it at the specified row.
    col_idx: 0-indexed column (0 = A, 1 = B, etc.)
    row_idx: 0-indexed row (0 = 1, 1 = 2, etc.)
    col_width_chars: character width set on the column (e.g. 18)
    row_offset_px: top padding in pixels
    """
    col_width_px = int(col_width_chars * 7.5 + 5)
    col_off_px = max(0, (col_width_px - img.width) // 2)

    marker = AnchorMarker(
        col=col_idx,
        colOff=pixels_to_EMU(col_off_px),
        row=row_idx,
        rowOff=pixels_to_EMU(row_offset_px),
    )
    size = XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height))
    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


def style_range(
    ws,
    cell_range: str,
    font: Optional[Font] = None,
    fill: Optional[PatternFill] = None,
    border: Optional[Border] = None,
    alignment: Optional[Alignment] = None,
):
    """
    Applies style attributes to all cells in a merged or rectangular range.
    Ensures complete visual consistency across Excel rendering engines.
    """
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if border is not None:
                cell.border = border
            if alignment is not None:
                cell.alignment = alignment
