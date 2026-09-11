from fastapi import APIRouter, Depends, Query, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import desc, func
from typing import List, Optional
from datetime import datetime, date, timezone
from pydantic import BaseModel
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database.connection import get_db
from app.database.models import CameraOfflineSession, MachineUtilization
from app.config import settings

router = APIRouter(prefix="/api/camera-sessions", tags=["camera-sessions"])

# --- Schemas ---

class OfflineSessionCreate(BaseModel):
    channel_key: str
    cam_ip: Optional[str] = None
    went_offline_at: datetime
    came_online_at: Optional[datetime] = None
    duration_seconds: Optional[float] = None

class OfflineSessionResponse(BaseModel):
    id: int
    mc_id: str
    channel_key: str
    cam_ip: Optional[str]
    date: date
    went_offline_at: datetime
    came_online_at: Optional[datetime]
    duration_seconds: Optional[float]

    class Config:
        from_attributes = True

# --- Endpoints ---

@router.post("/offline", summary="Record or update an offline session")
def record_offline_session(payload: OfflineSessionCreate, db: Session = Depends(get_db)):
    """
    Called by the factory_analytics script when a camera disconnects or reconnects.
    If came_online_at is None, a new session is created.
    If came_online_at is provided, it finds the latest open session for this channel and closes it.
    Since one channel can map to multiple machines (via roi_config.json), we duplicate the session per machine.
    """
    import os
    import json
    
    # Load ROI config to find machines for this channel
    roi_config_path = os.path.join(os.path.dirname(__file__), "..", "..", "factory_analytics", "roi_config.json")
    machines_for_channel = []
    if os.path.exists(roi_config_path):
        try:
            with open(roi_config_path, "r") as f:
                roi_data = json.load(f)
                if payload.channel_key in roi_data:
                    for roi in roi_data[payload.channel_key]:
                        if "machine_id" in roi and roi["machine_id"] not in machines_for_channel:
                            machines_for_channel.append(roi["machine_id"])
        except Exception:
            pass
            
    if not machines_for_channel:
        return {"status": "ignored", "reason": "No machines mapped to channel"}

    updated_sessions = []
    
    for mc_id in machines_for_channel:
        if payload.came_online_at is None:
            # Create new open session
            new_session = CameraOfflineSession(
                mc_id=mc_id,
                channel_key=payload.channel_key,
                cam_ip=payload.cam_ip,
                date=payload.went_offline_at.date(),
                went_offline_at=payload.went_offline_at
            )
            db.add(new_session)
            updated_sessions.append(new_session)
        else:
            # Close existing open session
            session = db.query(CameraOfflineSession).filter(
                CameraOfflineSession.mc_id == mc_id,
                CameraOfflineSession.channel_key == payload.channel_key,
                CameraOfflineSession.came_online_at == None
            ).order_by(desc(CameraOfflineSession.went_offline_at)).first()
            
            if session:
                session.came_online_at = payload.came_online_at
                if payload.duration_seconds is not None:
                    session.duration_seconds = payload.duration_seconds
                else:
                    session.duration_seconds = (payload.came_online_at - session.went_offline_at).total_seconds()
                updated_sessions.append(session)
            else:
                new_session = CameraOfflineSession(
                    mc_id=mc_id,
                    channel_key=payload.channel_key,
                    cam_ip=payload.cam_ip,
                    date=payload.went_offline_at.date(),
                    went_offline_at=payload.went_offline_at,
                    came_online_at=payload.came_online_at,
                    duration_seconds=payload.duration_seconds or (payload.came_online_at - payload.went_offline_at).total_seconds()
                )
                db.add(new_session)
                updated_sessions.append(new_session)
                
    db.commit()
    return {"status": "success", "sessions_updated": len(updated_sessions)}


@router.get("/report", summary="Get audit report for machines")
def get_audit_report(
    from_date: str = Query(..., alias="from", description="Start date YYYY-MM-DD"),
    to_date: str = Query(..., alias="to", description="End date YYYY-MM-DD"),
    mc_ids: Optional[str] = Query(None, description="Comma separated machine IDs"),
    db: Session = Depends(get_db)
):
    try:
        start_dt = datetime.strptime(from_date, "%Y-%m-%d").date()
        end_dt = datetime.strptime(to_date, "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date format. Use YYYY-MM-DD.")
        
    machine_list = [m.strip() for m in mc_ids.split(",")] if mc_ids and mc_ids.lower() != "all" else None
    
    # Fetch utilization data
    util_query = db.query(MachineUtilization).filter(
        MachineUtilization.date >= start_dt,
        MachineUtilization.date <= end_dt
    )
    if machine_list:
        util_query = util_query.filter(MachineUtilization.mc_id.in_(machine_list))
        
    util_rows = util_query.order_by(MachineUtilization.date.desc(), MachineUtilization.mc_id).all()
    
    # Fetch offline sessions
    session_query = db.query(CameraOfflineSession).filter(
        CameraOfflineSession.date >= start_dt,
        CameraOfflineSession.date <= end_dt
    )
    if machine_list:
        session_query = session_query.filter(CameraOfflineSession.mc_id.in_(machine_list))
        
    sessions = session_query.order_by(CameraOfflineSession.went_offline_at.desc()).all()
    
    # Group sessions by (date, mc_id)
    from collections import defaultdict
    session_map = defaultdict(list)
    for s in sessions:
        session_map[(s.date, s.mc_id)].append({
            "channel_key": s.channel_key,
            "went_offline_at": s.went_offline_at.isoformat() if s.went_offline_at else None,
            "came_online_at": s.came_online_at.isoformat() if s.came_online_at else None,
            "duration_seconds": s.duration_seconds
        })
        
    def _format_sec(sec):
        if sec is None: return "0h 0m 0s"
        total = int(round(sec))
        h = total // 3600
        m = (total % 3600) // 60
        s = total % 60
        return f"{h}h {m}m {s}s"

    report = []
    for row in util_rows:
        key = (row.date, row.mc_id)
        report.append({
            "date": str(row.date),
            "mc_id": row.mc_id,
            "runtime_seconds": row.runtime,
            "runtime_formatted": _format_sec(row.runtime),
            "downtime_seconds": row.downtime,
            "downtime_formatted": _format_sec(row.downtime),
            "offline_seconds": row.offline_time,
            "offline_formatted": _format_sec(row.offline_time),
            "total_seconds": row.total_time,
            "total_formatted": _format_sec(row.total_time),
            "utilization_percent": row.utilization_percent,
            "offline_sessions": session_map.get(key, [])
        })
        
    return report

@router.get("/report/download", summary="Download Audit Report as Excel")
def download_audit_report(
    from_date: str = Query(..., alias="from"),
    to_date: str = Query(..., alias="to"),
    mc_ids: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    report_data = get_audit_report(from_date, to_date, mc_ids, db)
    
    wb = openpyxl.Workbook()
    ws_summary = wb.active
    ws_summary.title = "Audit Summary"
    
    header_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(fill_type="solid", fgColor="1F3864")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )
    
    headers = [
        "Date", "Machine ID", "Runtime", "Downtime", "Offline Time", "Total Time", "Utilization %", "Offline Sessions"
    ]
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws_summary.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
    for row_idx, r in enumerate(report_data, start=2):
        values = [
            r["date"],
            r["mc_id"],
            r["runtime_formatted"],
            r["downtime_formatted"],
            r["offline_formatted"],
            r["total_formatted"],
            round(r["utilization_percent"], 2),
            len(r["offline_sessions"])
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws_summary.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = center_align
            
    for col_idx, header in enumerate(headers, start=1):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 4, 15)
        
    ws_details = wb.create_sheet(title="Offline Session Details")
    detail_headers = [
        "Date", "Machine ID", "Channel", "Went Offline At", "Came Online At", "Duration"
    ]
    
    for col_idx, header in enumerate(detail_headers, start=1):
        cell = ws_details.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = PatternFill(fill_type="solid", fgColor="B22222")
        cell.alignment = center_align
        cell.border = thin_border
        
    row_idx = 2
    for r in report_data:
        for s in r["offline_sessions"]:
            
            def _fmt_dt(dt_str):
                if not dt_str: return "Still Offline"
                try:
                    dt = datetime.fromisoformat(dt_str)
                    return dt.strftime("%Y-%m-%d %H:%M:%S")
                except:
                    return dt_str
                    
            dur_sec = s["duration_seconds"]
            if dur_sec is not None:
                total = int(round(dur_sec))
                h = total // 3600
                m = (total % 3600) // 60
                sec = total % 60
                dur_str = f"{h}h {m}m {sec}s"
            else:
                dur_str = "Ongoing"
                
            values = [
                r["date"],
                r["mc_id"],
                s["channel_key"],
                _fmt_dt(s["went_offline_at"]),
                _fmt_dt(s["came_online_at"]),
                dur_str
            ]
            for col_idx, val in enumerate(values, start=1):
                cell = ws_details.cell(row=row_idx, column=col_idx, value=val)
                cell.border = thin_border
                cell.alignment = center_align
            row_idx += 1
            
    for col_idx, header in enumerate(detail_headers, start=1):
        ws_details.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 4, 18)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    filename = f"Audit_Report_{from_date}_to_{to_date}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
