"""
Full pipeline: handwritten trade slip(s) -> structured extraction -> client-ready Word doc(s).

SETUP (once):
    pip install anthropic python-docx python-dotenv
    Create a .env file next to this script containing:
        ANTHROPIC_API_KEY=sk-ant-...

RUN (single image):
    python pipeline.py "path/to/image.jpg"

RUN (whole folder):
    python pipeline.py "path/to/folder"

RUN (folder + one combined summary doc for all of them):
    python pipeline.py "path/to/folder" --combine

Outputs, per image, land next to the source image:
    image.json   - raw structured extraction (for your own debugging)
    image.docx   - client-ready Word document

With --combine, an additional "Combined Report.docx" is created in the folder,
containing every processed slip as its own section.
"""

import argparse
import base64
import json
import sys
import time
from pathlib import Path

from anthropic import Anthropic
from dotenv import load_dotenv
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

load_dotenv()

IMAGE_EXTS = {".jpg", ".jpeg", ".png"}

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
ACCENT = RGBColor(0x2E, 0x6F, 0x7E)
RED = RGBColor(0xB9, 0x1C, 0x1C)
MUTED = RGBColor(0x6B, 0x72, 0x80)
INK = RGBColor(0x1A, 0x1A, 0x1A)

PROMPT = """You are reading a handwritten trade/inventory note. It may mix Bengali and
English, with measurements in feet/inches/mm, weights in kg/ton, piece counts ("pc"),
gauge specs, and vendor names.

Extract every distinct line item you can find as a JSON array. Each item should have:
- "dimension": the size/length spec as written (e.g. "5X8", "17 ft")
- "spec": any gauge/thickness/profile description
- "quantity": the quantity or weight value as written
- "unit": the unit (pc, kg, ton, etc.)
- "notes": anything else relevant (vendor name, date, price, short/balance notes)
- "confidence": "high" or "low" -- mark "low" if the handwriting is genuinely ambiguous

Also include one summary object at the end with "vendor" (business name on the page, if
visible) and "date" (if visible).

Respond with ONLY a valid JSON array. No prose, no markdown fences.
"""


# ---------- Extraction ----------

def extract_image(client: Anthropic, image_path: Path) -> list:
    media_type = "image/png" if image_path.suffix.lower() == ".png" else "image/jpeg"
    b64 = base64.b64encode(image_path.read_bytes()).decode("utf-8")

    resp = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=2000,
        messages=[
            {
                "role": "user",
                "content": [
                    {"type": "image", "source": {"type": "base64", "media_type": media_type, "data": b64}},
                    {"type": "text", "text": PROMPT},
                ],
            }
        ],
    )
    raw = resp.content[0].text.strip().replace("```json", "").replace("```", "").strip()
    return json.loads(raw)


# ---------- Word doc building ----------

def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=INK, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text if text not in (None, "") else "—")
    run.bold = bold
    run.font.color.rgb = color
    run.font.size = Pt(size)


def add_slip_section(doc: Document, items: list, source_name: str):
    """Add one slip's extraction as a titled table section into an existing Document."""
    summary = next((i for i in items if i.get("vendor") or (i.get("date") and not i.get("dimension"))), None)
    rows = [i for i in items if i.get("dimension")]

    title = doc.add_paragraph()
    run = title.add_run(summary.get("vendor") if summary and summary.get("vendor") else source_name)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = NAVY

    if summary and summary.get("date"):
        p = doc.add_paragraph()
        r = p.add_run("Date: " + summary["date"])
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED

    if not rows:
        p = doc.add_paragraph()
        r = p.add_run("No line items could be confidently extracted from this image.")
        r.italic = True
        r.font.color.rgb = MUTED
        return

    headers = ["Dimension", "Spec", "Quantity", "Unit", "Notes"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=11)
        shade_cell(cell, "2E6F7E")

    low_count = 0
    for idx, item in enumerate(rows):
        row = table.add_row()
        low_conf = item.get("confidence") == "low"
        if low_conf:
            low_count += 1
        fill = "FEF2F2" if low_conf else ("F4F6F8" if idx % 2 == 0 else "FFFFFF")
        color = RED if low_conf else INK
        notes = (item.get("notes") or "") + ("  (please verify)" if low_conf else "")
        values = [item.get("dimension"), item.get("spec"), item.get("quantity"), item.get("unit"), notes]
        for col, val in enumerate(values):
            set_cell_text(row.cells[col], val, color=color)
            shade_cell(row.cells[col], fill)

    if low_count > 0:
        p = doc.add_paragraph()
        r = p.add_run(
            f"Note: {low_count} item(s) highlighted in red had handwriting that was hard to "
            "read with certainty. Please double-check these against the original."
        )
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RED


def build_single_docx(items: list, out_path: Path, source_name: str):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    title = doc.add_paragraph()
    run = title.add_run("Extracted Order Details")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY
    doc.add_paragraph()
    add_slip_section(doc, items, source_name)
    doc.save(out_path)


# ---------- Pipeline ----------

def process_one(client: Anthropic, image_path: Path):
    print(f"  Reading {image_path.name} ...", end=" ", flush=True)
    try:
        items = extract_image(client, image_path)
    except json.JSONDecodeError:
        print("FAILED (model did not return clean JSON)")
        return None
    except Exception as e:
        print(f"FAILED ({e})")
        return None

    json_path = image_path.with_suffix(".json")
    json_path.write_text(json.dumps(items, indent=2, ensure_ascii=False), encoding="utf-8")

    docx_path = image_path.with_suffix(".docx")
    build_single_docx(items, docx_path, image_path.stem)

    print("done")
    return items


def main():
    parser = argparse.ArgumentParser(description="Handwritten slip -> structured Word doc pipeline")
    parser.add_argument("path", help="Path to a single image, or a folder of images")
    parser.add_argument("--combine", action="store_true", help="Also build one Combined Report.docx (folder mode only)")
    args = parser.parse_args()

    target = Path(args.path)
    if not target.exists():
        print(f"Path not found: {target}")
        sys.exit(1)

    client = Anthropic()

    if target.is_file():
        images = [target]
        folder = target.parent
    else:
        images = sorted(p for p in target.iterdir() if p.suffix.lower() in IMAGE_EXTS)
        folder = target
        if not images:
            print(f"No image files found in {folder}")
            sys.exit(1)

    print(f"Processing {len(images)} image(s)...")
    all_results = []  # list of (image_path, items)
    for img in images:
        items = process_one(client, img)
        if items:
            all_results.append((img, items))
        time.sleep(0.3)  # small pacing buffer between API calls

    print(f"\nDone: {len(all_results)}/{len(images)} succeeded.")

    if args.combine and target.is_dir() and all_results:
        combined = Document()
        section = combined.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        title = combined.add_paragraph()
        run = title.add_run("Combined Order Report")
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = NAVY
        combined.add_paragraph()

        for img, items in all_results:
            add_slip_section(combined, items, img.stem)
            combined.add_paragraph()  # spacer between sections

        out_path = folder / "Combined Report.docx"
        combined.save(out_path)
        print(f"Combined report saved: {out_path}")


if __name__ == "__main__":
    main()